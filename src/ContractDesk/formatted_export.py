"""Export a filtered, static copy of one source workbook, preserving its design.

This is part of the desktop application; original workbooks are never saved.
Formula results are copied from Excel's saved cache to avoid broken links after
other rows/sheets have been removed.
"""
from copy import copy, deepcopy

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidationList
from employee_names import canonical_employee


def _normalize(value):
    import unicodedata
    return unicodedata.normalize('NFD', str(value or '').lower().replace('đ', 'd')).encode('ascii', 'ignore').decode().strip()


def _translate(formula, old_cell, new_cell):
    if not formula or str(formula).startswith('"'):
        return formula
    value = str(formula)
    try:
        translated = Translator(value if value.startswith('=') else '=' + value, origin=old_cell).translate_formula(new_cell)
        return translated if value.startswith('=') else translated[1:]
    except (ValueError, TypeError):
        return formula


def write_formatted_export(records, destination, include_status=False):
    if not records:
        raise ValueError('Không có dòng để xuất.')
    source = records[0]['source']
    if any(r['source'] != source for r in records):
        raise ValueError('Mỗi tệp xuất phải dùng cùng một tệp gốc để giữ nguyên thiết kế.')
    book = openpyxl.load_workbook(source, rich_text=True, keep_links=False)
    cache = None
    try:
        cache = openpyxl.load_workbook(source, data_only=True, rich_text=True, keep_links=False)
        sheet = next((s for s in book if _normalize(s.title) == 'hop dong 2026'), book.worksheets[0])
        saved = cache[sheet.title]
        records = sorted(records, key=lambda r: (r['row'], r['data']['certificate']))
        width = sheet.max_column
        headers = {_normalize(c.value): c.column for c in sheet[1] if c.value is not None}
        certificate_col = headers.get('so chung thu / cvtt')
        certificate_date_col = headers.get('ngay chung thu')
        employee_col = headers.get('cv thuc hien')
        if certificate_col is None or certificate_date_col is None:
            raise ValueError('Không tìm thấy cột chứng thư/ngày chứng thư trong tệp gốc. Quét lại.')

        # Snapshot cells before clearing data. Cached values preserve formula results
        # without references to excluded contracts or external workbooks.
        snapshots = {}
        row_dimensions = {}
        for row_number in {1} | {r['row'] for r in records}:
            cells = []
            for cell in sheet[row_number][:width]:
                value = deepcopy(cell.value)
                if cell.data_type == 'f':
                    value = deepcopy(saved.cell(row_number, cell.column).value)
                    if value is None:
                        raise ValueError(f'Ô {cell.coordinate} chưa có kết quả công thức đã lưu. Mở Excel, tính lại và lưu tệp gốc trước khi xuất.')
                cells.append((value, copy(cell._style), copy(getattr(cell, 'comment', None)), copy(getattr(cell, 'hyperlink', None))))
            snapshots[row_number] = cells
            row_dimensions[row_number] = copy(sheet.row_dimensions[row_number])

        old_merges = list(sheet.merged_cells.ranges)
        for merge in old_merges:
            if merge.max_row > 1 and merge.max_row != merge.min_row:
                raise ValueError('Tệp gốc gộp ô qua nhiều dòng dữ liệu. Bỏ gộp các dòng hợp đồng trước khi xuất để không ghép nhầm thông tin.')
            sheet.unmerge_cells(str(merge))
        old_cf = deepcopy(sheet.conditional_formatting)
        old_dv = deepcopy(sheet.data_validations)
        old_tables = [deepcopy(table) for table in sheet.tables.values()]

        # Keep source theme, styles, columns, sheet/page settings and header/footer.
        for other in list(book.worksheets):
            if other is not sheet:
                book.remove(other)
        book.active = 0
        sheet.sheet_state = 'visible'
        book.defined_names.clear()  # Removed sheets/rows must not leave broken named ranges.
        sheet.delete_rows(1, sheet.max_row)
        sheet.row_dimensions.clear()
        sheet.tables.clear()
        sheet.conditional_formatting = ConditionalFormattingList()
        sheet.data_validations = DataValidationList()
        # Sorting/filter criteria from the source must not hide ready rows.
        sheet.auto_filter.filterColumn = []
        sheet.auto_filter.sortState = None

        mapping = [(1, 1)] + [(r['row'], target) for target, r in enumerate(records, 2)]
        for original_row, target_row in mapping:
            row_dim = copy(row_dimensions[original_row])
            row_dim.index = target_row
            row_dim.hidden = False
            sheet.row_dimensions[target_row] = row_dim
            for column, (value, style, comment, hyperlink) in enumerate(snapshots[original_row], 1):
                cell = sheet.cell(target_row, column)
                cell.value = deepcopy(value)
                cell._style = copy(style)
                cell.comment = copy(comment)
                if hyperlink:
                    cell.hyperlink = copy(hyperlink)
                    cell.hyperlink.ref = cell.coordinate
            for merge in old_merges:
                if merge.min_row == merge.max_row == original_row:
                    sheet.merge_cells(start_row=target_row, end_row=target_row, start_column=merge.min_col, end_column=merge.max_col)

        for target_row, record in enumerate(records, 2):
            original_row = record['row']
            if employee_col:
                sheet.cell(target_row, employee_col).value = canonical_employee(saved.cell(original_row, employee_col).value)
            # Split paired multiline certificates, keeping all other original columns.
            raw_certs = str(saved.cell(original_row, certificate_col).value or '')
            import re
            certs = [s.strip() for s in re.split(r'[\n;]+', raw_certs) if s.strip()]
            if len(certs) > 1:
                sheet.cell(target_row, certificate_col).value = record['data']['certificate']
                raw_dates = str(saved.cell(original_row, certificate_date_col).value or '')
                dates = [s.strip() for s in re.split(r'[\n;]+', raw_dates) if s.strip()]
                if len(certs) == len(dates) and record['data']['certificate'] in certs:
                    sheet.cell(target_row, certificate_date_col).value = dates[certs.index(record['data']['certificate'])]

        # Rebuild range-based formatting/validation at the new row locations.
        for original_row, target_row in mapping:
            for group in old_cf:
                for area in group.sqref.ranges:
                    if area.min_row <= original_row <= area.max_row:
                        old_anchor = f'{get_column_letter(area.min_col)}{area.min_row}'
                        new_anchor = f'{get_column_letter(area.min_col)}{target_row}'
                        target = f'{new_anchor}:{get_column_letter(area.max_col)}{target_row}'
                        for rule in old_cf[group]:
                            rule = deepcopy(rule)
                            rule.formula = [_translate(f, old_anchor, new_anchor) for f in rule.formula]
                            sheet.conditional_formatting.add(target, rule)
            for validation in old_dv.dataValidation:
                for area in validation.sqref.ranges:
                    if area.min_row <= original_row <= area.max_row:
                        rule = deepcopy(validation)
                        old_anchor = f'{get_column_letter(area.min_col)}{area.min_row}'
                        new_anchor = f'{get_column_letter(area.min_col)}{target_row}'
                        rule.sqref = f'{new_anchor}:{get_column_letter(area.max_col)}{target_row}'
                        rule.formula1 = _translate(rule.formula1, old_anchor, new_anchor)
                        rule.formula2 = _translate(rule.formula2, old_anchor, new_anchor)
                        sheet.add_data_validation(rule)

        last_row = len(records) + 1
        for table in old_tables:
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if min_row != 1:
                continue
            table.ref = f'{get_column_letter(min_col)}1:{get_column_letter(max_col)}{last_row}'
            table.totalsRowCount = 0
            table.totalsRowShown = False
            if table.autoFilter:
                table.autoFilter.ref = table.ref
                table.autoFilter.filterColumn = []
                table.autoFilter.sortState = None
            table.sortState = None
            for column in table.tableColumns:
                column.calculatedColumnFormula = None
                column.totalsRowFormula = None
                column.totalsRowFunction = None
                column.totalsRowLabel = None
            sheet.add_table(table)
        sheet.auto_filter.ref = None if sheet.tables else f'A1:{get_column_letter(width)}{last_row}'
        if sheet.print_area:
            sheet.print_area = f'A1:{get_column_letter(width + (2 if include_status else 0))}{last_row}'

        if include_status:
            header_template = max((c for c in sheet[1][:width] if c.value is not None), key=lambda c: c.column)
            for offset, label in enumerate(('TRẠNG THÁI THEO DÕI', 'LÝ DO CHƯA SẴN SÀNG'), 1):
                cell = sheet.cell(1, width + offset, label)
                cell._style = copy(header_template._style)
                cell.alignment = copy(cell.alignment)
                cell.alignment = openpyxl.styles.Alignment(horizontal=cell.alignment.horizontal, vertical='top', wrap_text=True)
                sheet.column_dimensions[get_column_letter(width + offset)].width = 30 if offset == 1 else 65
            for target_row, record in enumerate(records, 2):
                reason = record['detail'] or {
                    'uploaded': 'Đã tải lên; không tải lại.',
                    'changed': 'Thông tin thay đổi sau khi tải lên; cần đối chiếu website.',
                }.get(record['state'], record.get('reason', ''))
                for offset, value in enumerate((record['status'], reason), 1):
                    cell = sheet.cell(target_row, width + offset, value)
                    cell._style = copy(sheet.cell(target_row, header_template.column)._style)
                    cell.alignment = openpyxl.styles.Alignment(vertical='top', wrap_text=True)
            if not sheet.tables:
                sheet.auto_filter.ref = f'A1:{get_column_letter(width + 2)}{last_row}'
        # Reset scroll selection to a cell that exists in the exported sheet.
        for selection in sheet.sheet_view.selection:
            selection.activeCell = 'A1'
            selection.sqref = 'A1'
        book.save(destination)
    finally:
        if cache is not None:
            cache.close()
        book.close()
